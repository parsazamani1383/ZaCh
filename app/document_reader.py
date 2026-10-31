
from pathlib import Path


def read_txt_file(file_path: str) -> str:
    """
    Read a UTF-8 text file and return its contents.
    """
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"فایل پیدا نشد: {path.name}")

    if path.suffix.lower() != ".txt":
        raise ValueError("در این مرحله فقط فایل‌های TXT پشتیبانی می‌شوند.")

    # utf-8-sig also handles text files that begin with a UTF-8 BOM.
    try:
        return path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError(
            f"فایل «{path.name}» با کدگذاری UTF-8 قابل خواندن نیست."
        ) from exc

from pypdf import PdfReader


def read_pdf_file(file_path: str) -> str:
    """Extract text from a text-based PDF file."""
    try:
        reader = PdfReader(file_path)
        pages_text = []

        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages_text.append(text)

        return "\n\n".join(pages_text)

    except Exception as error:
        raise RuntimeError(f"خطا در خواندن فایل PDF: {error}") from error

from docx import Document
from zipfile import ZipFile
from lxml import etree


def read_docx_file(file_path: str) -> str:
    """Extract text from Word paragraphs and text boxes."""
    try:
        document = Document(file_path)

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Read text inside text boxes from the DOCX XML.
        with ZipFile(file_path, "r") as docx_zip:
            xml_content = docx_zip.read("word/document.xml")

        root = etree.fromstring(xml_content)
        namespaces = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        }

        text_boxes = root.xpath(
            ".//w:txbxContent//w:p",
            namespaces=namespaces,
        )

        for paragraph in text_boxes:
            text = "".join(
                paragraph.xpath(".//w:t/text()", namespaces=namespaces)
            ).strip()

            if text:
                paragraphs.append(text)

        return "\n".join(paragraphs)

    except Exception as error:
        raise RuntimeError(
            f"خطا در خواندن فایل Word: {error}"
        ) from error


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _format_cell(value) -> str:
    """Convert one spreadsheet cell to clean text."""
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    if hasattr(value, "isoformat"):
        # datetime / date / time
        try:
            if getattr(value, "hour", 0) == 0 and getattr(value, "minute", 0) == 0 \
                    and getattr(value, "second", 0) == 0 and hasattr(value, "date"):
                return value.date().isoformat()
        except Exception:
            pass
        return value.isoformat(sep=" ")

    return str(value).strip().replace("\n", " ")


def _rows_to_text(sheet_name, rows) -> str:
    lines = []

    for row in rows:
        cells = [_format_cell(cell) for cell in row]

        # Drop trailing empty cells and fully empty rows.
        while cells and not cells[-1]:
            cells.pop()

        if any(cells):
            lines.append(" | ".join(cells))

    if not lines:
        return ""

    return f"### شیت: {sheet_name}\n" + "\n".join(lines)


def read_excel_file(file_path: str) -> str:
    """Extract the text/values of every sheet in an Excel workbook.

    .xlsx / .xlsm are read with openpyxl (cached formula results are used),
    legacy .xls files are read with xlrd when it is installed.
    """
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"فایل پیدا نشد: {path.name}")

    extension = path.suffix.lower()

    try:
        sections = []

        if extension in (".xlsx", ".xlsm"):
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError(
                    "کتابخانه openpyxl نصب نیست. در همان محیطی که برنامه "
                    "اجرا می‌شود این دستور را بزن: pip install openpyxl"
                ) from exc

            def extract(read_only):
                workbook = load_workbook(
                    filename=str(path),
                    read_only=read_only,
                    data_only=True,
                )
                found = []
                try:
                    for sheet in workbook.worksheets:
                        text = _rows_to_text(
                            sheet.title,
                            sheet.iter_rows(values_only=True),
                        )
                        if text:
                            found.append(text)
                finally:
                    workbook.close()
                return found

            try:
                sections = extract(read_only=True)
            except Exception:
                # Some files (odd metadata) fail in streaming mode.
                sections = extract(read_only=False)

        elif extension == ".xls":
            try:
                import xlrd
            except ImportError as exc:
                raise RuntimeError(
                    "برای خواندن فایل‌های قدیمی XLS باید کتابخانه xlrd نصب شود "
                    "(pip install xlrd)؛ یا فایل را به‌صورت XLSX ذخیره کن."
                ) from exc

            workbook = xlrd.open_workbook(str(path))

            for sheet in workbook.sheets():
                rows = (
                    sheet.row_values(index)
                    for index in range(sheet.nrows)
                )
                text = _rows_to_text(sheet.name, rows)
                if text:
                    sections.append(text)

        else:
            raise ValueError("فرمت اکسل پشتیبانی نمی‌شود.")

        return "\n\n".join(sections)

    except (RuntimeError, ValueError):
        raise
    except Exception as error:
        raise RuntimeError(f"خطا در خواندن فایل Excel: {error}") from error


# ---------------------------------------------------------------------------
# Single entry point used by the UI
# ---------------------------------------------------------------------------

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".xlsx", ".xlsm", ".xls"}

FILE_DIALOG_FILTER = (
    "اسناد پشتیبانی‌شده (*.pdf *.docx *.txt *.xlsx *.xlsm *.xls);;"
    "PDF (*.pdf);;Word (*.docx);;Excel (*.xlsx *.xlsm *.xls);;TXT (*.txt)"
)


def read_document(file_path: str) -> str:
    """Read any supported document and return its text."""
    extension = Path(file_path).suffix.lower()

    if extension == ".txt":
        return read_txt_file(file_path)
    if extension == ".pdf":
        return read_pdf_file(file_path)
    if extension == ".docx":
        return read_docx_file(file_path)
    if extension in (".xlsx", ".xlsm", ".xls"):
        return read_excel_file(file_path)

    raise ValueError("فرمت این فایل پشتیبانی نمی‌شود.")
